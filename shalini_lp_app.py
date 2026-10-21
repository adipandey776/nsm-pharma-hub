import streamlit as st
import sqlite3
import pandas as pd
import openpyxl
from datetime import datetime, date
import re
import os
import io
import zipfile
from PIL import Image

try:
    import pytesseract
    tesseract_default = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.path.exists(tesseract_default):
        pytesseract.pytesseract.tesseract_cmd = tesseract_default
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

st.set_page_config(page_title="New Shalini Medical - LP Suite", layout="wide")

DB_FILE = "shalini_medical.db"
TEMPLATE_FILE = "BILL NO 2842.xlsx"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS inventory (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    product_name TEXT,
                    manufacturer TEXT,
                    batch_no TEXT,
                    exp_date TEXT,
                    category TEXT,
                    mrp_per_unit REAL,
                    purchase_cost_per_unit REAL,
                    stock_qty INTEGER DEFAULT 0
                )''')
    c.execute('''CREATE TABLE IF NOT EXISTS invoices (
                    invoice_no INTEGER PRIMARY KEY,
                    bill_date TEXT,
                    patient_name TEXT,
                    prescribed_by TEXT,
                    category_type TEXT,
                    gross_amount REAL,
                    gst_amount REAL,
                    discount_pct REAL,
                    discount_amount REAL,
                    net_billed_amount REAL,
                    total_purchase_cost REAL,
                    net_profit REAL,
                    payment_status TEXT DEFAULT 'Pending'
                )''')
    c.execute('''CREATE TABLE IF NOT EXISTS invoice_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    invoice_no INTEGER,
                    qty INTEGER,
                    product_name TEXT,
                    manufacturer TEXT,
                    batch_no TEXT,
                    exp_date TEXT,
                    mrp_amount REAL,
                    purchase_amount REAL
                )''')
    c.execute('''CREATE TABLE IF NOT EXISTS requisitions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    req_date TEXT,
                    patient_or_ward TEXT,
                    prescribed_by TEXT,
                    items_requested TEXT,
                    status TEXT DEFAULT 'Pending Supply'
                )''')
    c.execute("SELECT COUNT(*) FROM inventory")
    if c.fetchone()[0] == 0:
        c.execute('''INSERT INTO inventory 
                     (product_name, manufacturer, batch_no, exp_date, category, mrp_per_unit, purchase_cost_per_unit, stock_qty)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                  ("BISOCODYL SUPP", "NEON", "BLNP758007", "2027-01-01", "Medicine (5% Disc)", 37.80, 28.50, 100))
    conn.commit()
    conn.close()

init_db()

def extract_medicine_details(ocr_text):
    batch, exp, mrp = "", "", 0.0
    batch_match = re.search(r'(?:B\.?\s*No\.?|Batch\s*No\.?|Batch|LOT)[:\s\-]*([A-Z0-9]{4,14})', ocr_text, re.IGNORECASE)
    if batch_match:
        batch = batch_match.group(1).upper()
    exp_match = re.search(r'(?:Exp\.?|Expiry|Use\s*Before)[:\s\-]*((?:0[1-9]|1[0-2])[\/\-]\d{2,4}|(?:JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[A-Z]*[\s\-\/]*\d{2,4})', ocr_text, re.IGNORECASE)
    if exp_match:
        exp = exp_match.group(1).upper()
    mrp_match = re.search(r'(?:M\.?R\.?P\.?|Price)[:\s₹Rs\.]*(\d+(?:\.\d{1,2})?)', ocr_text, re.IGNORECASE)
    if mrp_match:
        try:
            mrp = float(mrp_match.group(1))
        except ValueError:
            mrp = 0.0
    return batch, exp, mrp

def generate_excel_invoice(invoice_no, bill_date, patient_name, doctor_name, discount_pct, items):
    os.makedirs("generated_bills", exist_ok=True)
    out_path = os.path.join("generated_bills", f"BILL NO {invoice_no}.xlsx")
    wb = openpyxl.load_workbook(TEMPLATE_FILE)
    sheet = wb.active
    sheet['A5'] = f"PATIENT's NAME:  {patient_name.upper()}"
    sheet['E5'] = datetime.strptime(str(bill_date), "%Y-%m-%d")
    sheet['A6'] = f"PRESCRIBED BY DR.: {doctor_name.upper()}"
    sheet['E6'] = int(invoice_no)
    for idx in range(11):
        row_num = 9 + idx
        if idx < len(items):
            item = items[idx]
            sheet.cell(row_num, 1, item['qty'])
            sheet.cell(row_num, 2, item['product_name'])
            sheet.cell(row_num, 3, item['manufacturer'])
            sheet.cell(row_num, 4, item['batch_no'])
            try:
                exp_dt = datetime.strptime(item['exp_date'], "%Y-%m-%d")
                sheet.cell(row_num, 5, exp_dt)
            except Exception:
                sheet.cell(row_num, 5, item['exp_date'])
            sheet.cell(row_num, 6, f"=G{row_num}*5%")
            sheet.cell(row_num, 7, round(item['mrp_amount'], 2))
        else:
            for col in [1, 2, 3, 4, 5, 7]:
                sheet.cell(row_num, col, None)
            sheet.cell(row_num, 6, f"=G{row_num}*5%")
    sheet['G20'] = "=SUM(G9:G19)"
    sheet['G21'] = "=SUM(F9:F19)"
    sheet['F22'] = discount_pct / 100.0
    sheet['G22'] = f"=G20*{int(discount_pct)}/100"
    sheet['G23'] = "=G20-G22"
    wb.save(out_path)
    return out_path

st.title("🏥 New Shalini Medical — Railway LP Billing & Profit Suite")
st.caption("Rangatand, Dhanbad | GSTIN: 20ANNPP3569M1ZH | Drug Lic. No.: 147415(20)")

tab1, tab2, tab3, tab4 = st.tabs([
    "🧾 1. Create Weekly Bill", 
    "📋 2. Hospital Requisitions (Indents)",
    "📸 3. Scan / Add Medicine & Stock", 
    "📊 4. Monthly Profit, Top-Sheet & Payments"
])

with tab1:
    conn = sqlite3.connect(DB_FILE)
    inv_df = pd.read_sql_query("SELECT * FROM inventory ORDER BY product_name", conn)
    c = conn.cursor()
    c.execute("SELECT MAX(invoice_no) FROM invoices")
    max_inv = c.fetchone()[0]
    next_invoice_no = (max_inv + 1) if max_inv else 2843
    conn.close()

    default_patient = st.session_state.get("prefill_patient", "UDAY KUMAR")
    default_doc = st.session_state.get("prefill_doc", "CMS DHN")

    col1, col2, col3, col4 = st.columns(4)
    with col1:
        bill_no = st.number_input("Starting Invoice No.", value=next_invoice_no, step=1)
    with col2:
        bill_date = st.date_input("Bill Date", value=date.today())
    with col3:
        patient_name = st.text_input("Patient / Ward / Officer Name", value=default_patient)
    with col4:
        prescribed_by = st.text_input("Prescribed By Dr.", value=default_doc)

    bill_type = st.radio(
        "Bill Category (Sets Discount Automatically)", 
        ["Medicine (5% Discount)", "Surgical Item (10% Discount)"], 
        horizontal=True
    )
    disc_pct = 5.0 if "5%" in bill_type else 10.0

    st.markdown("#### Add Items to Bill (Auto-splits into multiple bills if >11 items)")
    if "cart" not in st.session_state:
        st.session_state.cart = []

    if not inv_df.empty:
        c1, c2, c3 = st.columns([3, 1, 1])
        with c1:
            inv_df["display"] = (
                inv_df["product_name"] + " | Batch: " + inv_df["batch_no"] + 
                " | Exp: " + inv_df["exp_date"] + " | Unit MRP: ₹" + inv_df["mrp_per_unit"].astype(str) +
                " | Stock: " + inv_df["stock_qty"].astype(str)
            )
            selected_option = st.selectbox("Search & Select Medicine / Surgical Item", inv_df["display"].tolist())
        with c2:
            qty = st.number_input("Quantity (Units)", min_value=1, value=30, step=1)
        with c3:
            st.write("")
            st.write("")
            if st.button("➕ Add to Bill", use_container_width=True):
                row = inv_df[inv_df["display"] == selected_option].iloc[0]
                mrp_total = float(row["mrp_per_unit"]) * qty
                purchase_total = float(row["purchase_cost_per_unit"]) * qty
                st.session_state.cart.append({
                    "inv_id": int(row["id"]),
                    "qty": int(qty),
                    "product_name": row["product_name"],
                    "manufacturer": row["manufacturer"],
                    "batch_no": row["batch_no"],
                    "exp_date": row["exp_date"],
                    "unit_mrp": float(row["mrp_per_unit"]),
                    "GST 5%": round(mrp_total * 0.05, 2),
                    "mrp_amount": round(mrp_total, 2),
                    "purchase_amount": round(purchase_total, 2)
                })

    if st.session_state.cart:
        cart_df = pd.DataFrame(st.session_state.cart)
        st.dataframe(
            cart_df[["qty", "product_name", "manufacturer", "batch_no", "exp_date", "GST 5%", "mrp_amount"]],
            use_container_width=True
        )

        num_sheets = (len(st.session_state.cart) + 10) // 11
        if num_sheets > 1:
            st.info(f"⚡ You have {len(st.session_state.cart)} items. The app will automatically generate **{num_sheets} sequential bills** (Bill No. {bill_no} to {bill_no + num_sheets - 1}).")

        gross_total = cart_df["mrp_amount"].sum()
        gst_total = cart_df["GST 5%"].sum()
        disc_amount = round(gross_total * (disc_pct / 100.0), 2)
        net_total = round(gross_total - disc_amount, 2)
        purchase_cost_total = cart_df["purchase_amount"].sum()
        estimated_profit = round(net_total - purchase_cost_total, 2)

        m1, m2, m3, m4, m5 = st.columns(5)
        m1.metric("TOTAL AMOUNT (Gross)", f"₹{gross_total:,.2f}")
        m2.metric("GST AMOUNT (5%)", f"₹{gst_total:,.2f}")
        m3.metric(f"DISCOUNT ({int(disc_pct)}%)", f"-₹{disc_amount:,.2f}")
        m4.metric("TOTAL RS. (Net Bill)", f"₹{net_total:,.2f}")
        m5.metric("Your Net Profit", f"₹{estimated_profit:,.2f}")

        b_col1, b_col2 = st.columns([1, 1])
        with b_col1:
            if st.button("✅ Save & Generate Excel Bill(s)", type="primary", use_container_width=True):
                conn = sqlite3.connect(DB_FILE)
                c = conn.cursor()
                generated_files = []
                for chunk_idx in range(num_sheets):
                    curr_inv_no = int(bill_no) + chunk_idx
                    chunk_items = st.session_state.cart[chunk_idx * 11 : (chunk_idx + 1) * 11]
                    c_gross = sum(x["mrp_amount"] for x in chunk_items)
                    c_gst = sum(x["GST 5%"] for x in chunk_items)
                    c_disc = round(c_gross * (disc_pct / 100.0), 2)
                    c_net = round(c_gross - c_disc, 2)
                    c_purch = sum(x["purchase_amount"] for x in chunk_items)
                    c_prof = round(c_net - c_purch, 2)

                    c.execute('''INSERT OR REPLACE INTO invoices VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                              (curr_inv_no, str(bill_date), patient_name.upper(), prescribed_by.upper(),
                               bill_type, c_gross, c_gst, disc_pct, c_disc, c_net,
                               c_purch, c_prof, "Pending"))
                    c.execute("DELETE FROM invoice_items WHERE invoice_no = ?", (curr_inv_no,))
                    for item in chunk_items:
                        c.execute('''INSERT INTO invoice_items 
                                     (invoice_no, qty, product_name, manufacturer, batch_no, exp_date, mrp_amount, purchase_amount)
                                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                                  (curr_inv_no, item["qty"], item["product_name"], item["manufacturer"],
                                   item["batch_no"], item["exp_date"], item["mrp_amount"], item["purchase_amount"]))
                        c.execute("UPDATE inventory SET stock_qty = MAX(0, stock_qty - ?) WHERE id = ?", (item["qty"], item["inv_id"]))

                    if os.path.exists(TEMPLATE_FILE):
                        excel_path = generate_excel_invoice(
                            curr_inv_no, bill_date, patient_name, prescribed_by, disc_pct, chunk_items
                        )
                        generated_files.append(excel_path)

                conn.commit()
                conn.close()

                if len(generated_files) == 1:
                    with open(generated_files[0], "rb") as f:
                        st.download_button(
                            label=f"📥 Download BILL NO {bill_no}.xlsx",
                            data=f,
                            file_name=f"BILL NO {bill_no}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True
                        )
                elif len(generated_files) > 1:
                    zip_buffer = io.BytesIO()
                    with zipfile.ZipFile(zip_buffer, "w") as zf:
                        for fpath in generated_files:
                            zf.write(fpath, os.path.basename(fpath))
                    st.download_button(
                        label=f"📥 Download All {len(generated_files)} Bills (ZIP)",
                        data=zip_buffer.getvalue(),
                        file_name=f"BILLS_{bill_no}_TO_{bill_no + len(generated_files) - 1}.zip",
                        mime="application/zip",
                        use_container_width=True
                    )
                st.success("Saved to database, updated stock, and generated Excel bill(s)!")

        with b_col2:
            if st.button("🗑️ Clear Current Items", use_container_width=True):
                st.session_state.cart = []
                st.rerun()

with tab2:
    st.subheader("📋 Monthly / Daily Hospital Requisition Tracker")
    with st.form("req_form"):
        r1, r2, r3 = st.columns([1, 1, 2])
        with r1:
            req_dt = st.date_input("Indent Date", value=date.today())
            req_pat = st.text_input("Ward / Officer / Patient Name*", placeholder="e.g. MALE MEDICAL WARD").upper()
        with r2:
            req_doc = st.text_input("Prescribed By", value="CMS DHN").upper()
            req_stat = st.selectbox("Status", ["Pending Supply", "Supplied - Ready to Bill", "Billed"])
        with r3:
            req_items = st.text_area("Medicines / Surgical Items Needed*", placeholder="e.g. 30x Bisocodyl Supp...")
        if st.form_submit_button("➕ Save Requisition", type="primary"):
            if req_pat and req_items:
                conn = sqlite3.connect(DB_FILE)
                conn.execute("INSERT INTO requisitions (req_date, patient_or_ward, prescribed_by, items_requested, status) VALUES (?, ?, ?, ?, ?)",
                             (str(req_dt), req_pat, req_doc, req_items, req_stat))
                conn.commit()
                conn.close()
                st.success("Requisition saved!")
                st.rerun()

    conn = sqlite3.connect(DB_FILE)
    req_df = pd.read_sql_query("SELECT * FROM requisitions ORDER BY id DESC", conn)
    conn.close()

    if not req_df.empty:
        for _, r_row in req_df.iterrows():
            with st.expander(f"📌 [{r_row['status']}] {r_row['patient_or_ward']} ({r_row['req_date']}) — {r_row['items_requested'][:50]}..."):
                st.write(f"**Prescribed By:** {r_row['prescribed_by']}")
                st.write(f"**Items Needed:** {r_row['items_requested']}")
                col_a, col_b, col_c = st.columns(3)
                if col_a.button("🧾 Load Patient into Bill Creator", key=f"load_{r_row['id']}"):
                    st.session_state["prefill_patient"] = r_row["patient_or_ward"]
                    st.session_state["prefill_doc"] = r_row["prescribed_by"]
                    st.success(f"Loaded {r_row['patient_or_ward']} into Tab 1!")
                if col_b.button("✅ Mark as Billed", key=f"done_{r_row['id']}"):
                    conn = sqlite3.connect(DB_FILE)
                    conn.execute("UPDATE requisitions SET status = 'Billed' WHERE id = ?", (r_row['id'],))
                    conn.commit()
                    conn.close()
                    st.rerun()
                if col_c.button("🗑️ Delete", key=f"del_{r_row['id']}"):
                    conn = sqlite3.connect(DB_FILE)
                    conn.execute("DELETE FROM requisitions WHERE id = ?", (r_row['id'],))
                    conn.commit()
                    conn.close()
                    st.rerun()

with tab3:
    st.subheader("📸 Upload Medicine Photo to Auto-Fill Batch, Exp & MRP")
    uploaded_img = st.file_uploader("Upload or take a photo of medicine strip/box", type=["jpg", "jpeg", "png"])
    ocr_batch, ocr_exp, ocr_mrp = "", "2027-01-01", 0.0
    if uploaded_img is not None:
        img = Image.open(uploaded_img)
        st.image(img, caption="Uploaded Strip/Box", width=280)
        if OCR_AVAILABLE:
            try:
                raw_text = pytesseract.image_to_string(img)
                b_ext, e_ext, m_ext = extract_medicine_details(raw_text)
                if b_ext: ocr_batch = b_ext
                if e_ext: ocr_exp = e_ext
                if m_ext: ocr_mrp = m_ext
                st.info("OCR extracted values populated below!")
            except Exception as e:
                st.warning(f"OCR notice: {e}. Enter details manually below.")

    st.markdown("#### Save Medicine / Surgical Item to Memory")
    with st.form("add_inventory_form"):
        f1, f2, f3 = st.columns(3)
        with f1:
            p_name = st.text_input("Brand / Product Name*", placeholder="e.g. BISOCODYL SUPP").upper()
            mfg = st.text_input("Manufacturer*", placeholder="e.g. NEON").upper()
            cat = st.selectbox("Category", ["Medicine (5% Disc)", "Surgical (10% Disc)"])
        with f2:
            b_no = st.text_input("Batch No.*", value=ocr_batch, placeholder="e.g. BLNP758007").upper()
            exp_dt = st.text_input("Expiry Date (YYYY-MM-DD)*", value=ocr_exp, placeholder="2027-01-01")
            qty_bought = st.number_input("Total Units Purchased for Stock", min_value=0, value=100, step=10)
        with f3:
            pack_size = st.number_input("Units per Strip/Pack (Use 1 if priced per single unit)", min_value=1, value=1)
            pack_mrp = st.number_input("Printed MRP on Pack (₹)*", min_value=0.0, value=float(ocr_mrp), step=0.5)
            pack_cost = st.number_input("Your Distributor Purchase Price for Pack (₹)*", min_value=0.0, value=0.0, step=0.5)

        if st.form_submit_button("💾 Save Item to Inventory Memory", type="primary") and p_name and b_no:
            unit_mrp = round(pack_mrp / pack_size, 4)
            unit_cost = round(pack_cost / pack_size, 4)
            conn = sqlite3.connect(DB_FILE)
            conn.execute('''INSERT INTO inventory 
                            (product_name, manufacturer, batch_no, exp_date, category, mrp_per_unit, purchase_cost_per_unit, stock_qty)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                         (p_name, mfg, b_no, exp_dt, cat, unit_mrp, unit_cost, int(qty_bought)))
            conn.commit()
            conn.close()
            st.success(f"Saved {p_name} (Batch: {b_no}) | Unit MRP: ₹{unit_mrp:.2f} | Unit Cost: ₹{unit_cost:.2f}")
            st.rerun()

    st.markdown("#### Saved Inventory & Live Stock")
    conn = sqlite3.connect(DB_FILE)
    saved_df = pd.read_sql_query("SELECT * FROM inventory ORDER BY id DESC", conn)
    conn.close()
    st.dataframe(saved_df, use_container_width=True)

with tab4:
    st.subheader("📊 Monthly Distributor Purchase vs. Railway Billing Profit")
    conn = sqlite3.connect(DB_FILE)
    bills_df = pd.read_sql_query("SELECT * FROM invoices ORDER BY invoice_no DESC", conn)
    conn.close()

    if bills_df.empty:
        st.info("No bills generated yet. Create your first bill in Tab 1!")
    else:
        bills_df["Month"] = bills_df["bill_date"].str.slice(0, 7)
        months = sorted(bills_df["Month"].unique().tolist(), reverse=True)
        selected_month = st.selectbox("Select Month", ["All", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"])
